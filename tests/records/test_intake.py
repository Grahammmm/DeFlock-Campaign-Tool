import hashlib, importlib.util, io, json, os, pathlib, sqlite3, subprocess, sys, tempfile, unittest, zipfile
from email.message import EmailMessage

SCRIPT=pathlib.Path(__file__).resolve().parents[2]/"campaign_tool"/"records"/"intake"/"folder.py"
spec=importlib.util.spec_from_file_location("intake",SCRIPT); m=importlib.util.module_from_spec(spec); spec.loader.exec_module(m)

class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.base=pathlib.Path(self.tmp.name); self.root=self.base/"originals"; self.root.mkdir(); self.out=self.base/"out"
    def tearDown(self): self.tmp.cleanup()
    def run_cli(self,command="run"):
        p=subprocess.run([sys.executable,str(SCRIPT),command,"--root",str(self.root),"--output",str(self.out)],capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stderr); return p
    def db(self):
        db=sqlite3.connect(self.out/"intake.sqlite"); db.row_factory=sqlite3.Row; self.addCleanup(db.close); return db
    def test_duplicate_occurrences_idempotence_and_changed_bytes(self):
        (self.root/"a.txt").write_text("same\n"); (self.root/"b.txt").write_text("same\n"); self.run_cli()
        db=self.db(); self.assertEqual(db.execute("SELECT count(*) FROM docs").fetchone()[0],1)
        before=[tuple(r) for r in db.execute("SELECT oid,sha,receipt,first_seen FROM occurrences ORDER BY oid")]
        digest=db.execute("SELECT digest FROM docs").fetchone()[0]; self.run_cli()
        self.assertEqual(before,[tuple(r) for r in db.execute("SELECT oid,sha,receipt,first_seen FROM occurrences ORDER BY oid")])
        self.assertEqual(digest,db.execute("SELECT digest FROM docs").fetchone()[0])
        self.assertEqual(db.execute("SELECT count(*) FROM units").fetchone()[0],1)
        st=(self.root/"a.txt").stat(); (self.root/"a.txt").write_text("diff\n"); os.utime(self.root/"a.txt",ns=(st.st_atime_ns,st.st_mtime_ns)); self.run_cli()
        self.assertEqual(db.execute("SELECT count(*) FROM docs").fetchone()[0],2)
        self.assertEqual(db.execute("SELECT count(*) FROM occurrences").fetchone()[0],3)
    def test_malformed_pdf_visible(self):
        (self.root/"bad.pdf").write_bytes(b"%PDF-1.7\ninvalid"); self.run_cli()
        d=self.db().execute("SELECT * FROM docs").fetchone(); self.assertEqual(d["stage"],"failed"); self.assertEqual(d["review_status"],"not_reviewed")
        self.assertTrue(json.loads(d["digest"])["issues"])
    def test_csv_whitespace_redaction_duplicate_and_multiline_locators(self):
        (self.root/"log.csv").write_text('reason,user\n"   ",a\n[REDACTED],b\nvalid,c\nvalid,c\n"two\nlines",d\nshort\n')
        self.run_cli(); db=self.db(); d=json.loads(db.execute("SELECT digest FROM docs").fetchone()[0])
        self.assertEqual(d["counts"]["csv_row"],6); self.assertEqual(d["counts"]["duplicate_rows"],1)
        counts=d["counts"]["fields"][0]["counts"]; self.assertEqual(counts["whitespace"],1); self.assertEqual(counts["redacted_marker"],1)
        u=db.execute("SELECT locator FROM units WHERE kind='csv_row' AND ordinal=6").fetchone(); self.assertEqual(json.loads(u[0]),dict(record=6,line_start=6,line_end=7))
        self.assertEqual(d["counts"]["fields"][1]["counts"]["missing_field"],1)
    def test_zip_traversal_nested_duplicates_and_secrets(self):
        nested=io.BytesIO()
        with zipfile.ZipFile(nested,"w") as z: z.writestr("ok.txt","record")
        with zipfile.ZipFile(self.root/"pack.zip","w") as z:
            z.writestr("../escape.txt","no"); z.writestr("/absolute.txt","no"); z.writestr("C:\\escape.txt","no"); z.writestr("a\\..\\escape.txt","no")
            z.writestr(".env","do not read"); z.writestr("one.txt","record"); z.writestr("two.txt","record"); z.writestr("nested.zip",nested.getvalue())
            info=zipfile.ZipInfo("link"); info.create_system=3; info.external_attr=(0o120777<<16); z.writestr(info,"/etc/passwd")
        self.run_cli(); db=self.db()
        self.assertFalse((self.base/"escape.txt").exists()); self.assertEqual(db.execute("SELECT count(*) FROM edges").fetchone()[0],4)
        codes=[r[0] for r in db.execute("SELECT code FROM events")]; self.assertEqual(codes.count("unsafe_zip_member"),5)
        self.assertIn("excluded_credential_or_secret_name",codes)
        self.assertEqual(db.execute("SELECT count(*) FROM docs WHERE format='txt'").fetchone()[0],1)
    def test_eml_parent_child_and_headers(self):
        msg=EmailMessage(); msg["From"]="sender@example.invalid"; msg["Subject"]="fixture"; msg.set_content("Body text")
        msg.add_attachment(b"reason\nvalid\n",maintype="text",subtype="csv",filename="log.csv")
        (self.root/"mail.eml").write_bytes(msg.as_bytes()); self.run_cli(); db=self.db()
        self.assertEqual(db.execute("SELECT count(*) FROM edges").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM units WHERE kind='email_headers'").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM units WHERE kind='csv_row'").fetchone()[0],1)
    def test_symlinks_excluded(self):
        (self.root/"secret.key").write_text("secret"); (self.root/"link.txt").symlink_to(self.root/"secret.key"); self.run_cli("inventory")
        s=json.loads((self.out/"inventory-summary.json").read_text()); self.assertEqual(s["counts"].get("included",0),0); self.assertEqual(s["counts"]["excluded"],2)
    def test_pdf_all_blank_pages_queued_for_ocr(self):
        from pypdf import PdfWriter
        w=PdfWriter(); w.add_blank_page(width=72,height=72); w.add_blank_page(width=72,height=72)
        with (self.root/"blank.pdf").open("wb") as f: w.write(f)
        self.run_cli(); db=self.db(); d=json.loads(db.execute("SELECT digest FROM docs").fetchone()[0])
        self.assertEqual(d["counts"]["pages_expected"],2); self.assertEqual(d["counts"]["pdf_page"],2); self.assertEqual(d["counts"]["ocr_needed"],2)
    def test_xlsx_hidden_sheet_and_formula(self):
        from openpyxl import Workbook
        w=Workbook(); w.active["A1"]="visible"; s=w.create_sheet("Hidden"); s.sheet_state="hidden"; s["B2"]="=1+1"; w.save(self.root/"book.xlsx")
        self.run_cli(); db=self.db(); sheets=[json.loads(r[0]) for r in db.execute("SELECT data FROM units WHERE kind='workbook_sheet'")]
        self.assertEqual(len(sheets),2); self.assertEqual(sheets[1]["state"],"hidden")
        self.assertIn("=1+1",db.execute("SELECT group_concat(text) FROM units").fetchone()[0])
    def test_source_changed_after_inventory_uses_preserved_original(self):
        p=self.root/"a.txt"; p.write_text("old"); self.run_cli("inventory"); p.write_text("new"); self.run_cli("extract")
        db=self.db(); d=db.execute("SELECT sha,stage FROM docs").fetchone(); self.assertEqual(d[1],"complete")
        self.assertEqual((self.out/"blobs"/d[0]).read_text(),"old"); self.assertEqual(db.execute("SELECT text FROM units").fetchone()[0],"old")
        self.run_cli("inventory"); self.assertEqual(db.execute("SELECT count(*) FROM preservations").fetchone()[0],2)
        self.assertEqual((self.out/"blobs"/d[0]).stat().st_mode & 0o222,0)
    def test_original_directory_beats_audit_research_review_filename(self):
        (self.root/"audit-report.txt").write_text("agency audit"); (self.root/"research-review.txt").write_text("agency review"); self.run_cli()
        self.assertEqual({r[0] for r in self.db().execute("SELECT class FROM occurrences")},{"original"})
    def test_runtime_exclusions_and_extensionless_manual_queue(self):
        for name in ("tool.mjs","tool.cjs","api.h","lib.a","py.typed"): (self.root/name).write_text("toolchain")
        for name in ("site-packages","thing.dist-info","thing.egg-info"):
            p=self.root/name; p.mkdir(); (p/"METADATA").write_text("toolchain")
        (self.root/"opaque").write_bytes(b"unrecognized record"); self.run_cli()
        q=[json.loads(x) for x in (self.out/"all-agency-document-queue.jsonl").read_text().splitlines()]
        self.assertEqual(len(q),1); self.assertTrue(q[0]["name_review_required"]); self.assertEqual(q[0]["stage"],"unsupported")
    def test_extensionless_eml_body_attachment_and_xlsx_alias(self):
        from openpyxl import Workbook
        msg=EmailMessage(); msg["From"]="sender@example.invalid"; msg["To"]="recipient@example.invalid"; msg["Subject"]="fixture"; msg.set_content("Extensionless email body")
        msg.add_attachment(b"reason\nvalid\n",maintype="text",subtype="csv",filename="log.csv")
        raw=msg.as_bytes(); (self.root/hashlib.sha256(raw).hexdigest()).write_bytes(raw)
        w=Workbook(); w.active["A1"]="=2+3"; data=io.BytesIO(); w.save(data); raw=data.getvalue()
        (self.root/("000-"+hashlib.sha256(raw).hexdigest())).write_bytes(raw); (self.root/"z.xlsx").write_bytes(raw)
        self.run_cli(); db=self.db()
        self.assertEqual(db.execute("SELECT count(*) FROM docs WHERE format='eml'").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM docs WHERE format='xlsx'").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM units WHERE kind='email_body'").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM units WHERE kind='csv_row'").fetchone()[0],1)
        self.assertEqual(db.execute("SELECT count(*) FROM units WHERE kind='workbook_sheet'").fetchone()[0],1)
    def test_quarantine_excludes_aliases_without_touching_originals(self):
        personal=self.root/"personal.txt"; personal.write_text("synthetic unrelated fixture")
        (self.root/"alias.txt").write_bytes(personal.read_bytes()); (self.root/"flock.txt").write_text("in-scope fixture"); self.run_cli()
        p=subprocess.run([sys.executable,str(SCRIPT.parent/"scope_quarantine.py"),"--output",str(self.out),str(personal)],capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stderr); self.assertEqual(personal.read_text(),"synthetic unrelated fixture")
        self.run_cli(); q=[json.loads(x) for x in (self.out/"all-agency-document-queue.jsonl").read_text().splitlines()]
        self.assertEqual(len(q),1); self.assertEqual(self.db().execute("SELECT count(*) FROM units").fetchone()[0],1)
    def test_selected_retry_preserves_other_document_and_attempt_history(self):
        (self.root/"one.txt").write_text("one"); (self.root/"two.txt").write_text("two"); self.run_cli()
        one=hashlib.sha256(b"one").hexdigest(); two=hashlib.sha256(b"two").hexdigest(); db=self.db()
        old=db.execute("SELECT digest FROM docs WHERE sha=?",(two,)).fetchone()[0]
        p=subprocess.run([sys.executable,str(SCRIPT),"extract","--output",str(self.out),"--sha",one,"--retry"],capture_output=True,text=True)
        self.assertEqual(p.returncode,0,p.stderr); self.assertEqual(json.loads(p.stdout.splitlines()[-1])["processed"],1)
        self.assertEqual(old,db.execute("SELECT digest FROM docs WHERE sha=?",(two,)).fetchone()[0])
        attempts=list(db.execute("SELECT artifact_dir FROM extraction_attempts WHERE sha=?",(one,)))
        self.assertEqual(len(attempts),2); self.assertTrue(all(pathlib.Path(r[0],"digest.json").exists() for r in attempts))
    def test_invalid_unknown_and_inactive_selectors_block(self):
        p=self.root/"a.txt"; p.write_text("old"); self.run_cli(); old=hashlib.sha256(b"old").hexdigest(); p.write_text("new"); self.run_cli("inventory")
        db=self.db(); before=db.execute("SELECT count(*) FROM extraction_attempts").fetchone()[0]
        for sha in ("invalid","f"*64,old):
            result=subprocess.run([sys.executable,str(SCRIPT),"extract","--output",str(self.out),"--sha",sha,"--retry"],capture_output=True,text=True)
            self.assertNotEqual(result.returncode,0); self.assertIn("invalid/inactive selector rejected",result.stderr)
        self.assertEqual(before,db.execute("SELECT count(*) FROM extraction_attempts").fetchone()[0])
    def test_passwordless_encrypted_pdf_only(self):
        from pypdf import PdfWriter
        for name,password in (("open.pdf",""),("locked.pdf","synthetic-required")):
            w=PdfWriter(); w.add_blank_page(width=72,height=72); w.encrypt(user_password=password,owner_password="synthetic-owner")
            with (self.root/name).open("wb") as f: w.write(f)
        self.run_cli(); db=self.db()
        rows={pathlib.Path(r[0]).name:json.loads(r[1]) for r in db.execute("SELECT o.path,d.digest FROM occurrences o JOIN docs d ON d.sha=o.sha")}
        self.assertEqual(rows["open.pdf"]["counts"]["standard_passwordless_open"],1)
        self.assertEqual(rows["locked.pdf"]["stage"],"failed"); self.assertIn("requires_user_password",str(rows["locked.pdf"]["issues"]))
    def test_verified_legacy_blob_becomes_read_only_without_touching_source(self):
        source=self.root/"record.txt"; source.write_bytes(b"record"); mode=source.stat().st_mode; self.run_cli()
        sha=hashlib.sha256(b"record").hexdigest(); blob=self.out/"blobs"/sha; blob.chmod(0o600)
        pending=self.out/"blobs"/".pending-fixture"; pending.write_bytes(b"record"); m.seal_blob(pending,sha,self.out)
        self.assertEqual(blob.stat().st_mode & 0o777,0o400); self.assertEqual(blob.read_bytes(),b"record")
        self.assertEqual(source.stat().st_mode,mode); self.assertEqual(source.read_bytes(),b"record")
    def test_corrupt_blob_and_shared_original_inode_are_not_chmodded(self):
        source=self.root/"record.txt"; source.write_bytes(b"record"); self.run_cli()
        sha=hashlib.sha256(b"record").hexdigest(); blob=self.out/"blobs"/sha; blob.chmod(0o600); blob.write_bytes(b"corrupt")
        with self.assertRaisesRegex(ValueError,"hash_mismatch"): m.verify_blob(sha,self.out,True)
        self.assertEqual(blob.stat().st_mode & 0o777,0o600)
        blob.unlink(); os.link(source,blob); mode=source.stat().st_mode
        with self.assertRaisesRegex(ValueError,"shared_inode"): m.verify_blob(sha,self.out,True)
        self.assertEqual(source.stat().st_mode,mode)
    def test_docx_and_extensionless_alias_use_docx_handler_not_zip(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,"w") as z:
            z.writestr("[Content_Types].xml","<Types/>"); z.writestr("word/document.xml","<document/>")
        raw=data.getvalue(); (self.root/"response.docx").write_bytes(raw); (self.root/hashlib.sha256(raw).hexdigest()).write_bytes(raw); self.run_cli()
        db=self.db(); doc=db.execute("SELECT format,stage FROM docs").fetchone()
        self.assertEqual(tuple(doc),("docx","complete")); self.assertEqual(db.execute("SELECT count(*) FROM edges").fetchone()[0],0)
        self.assertEqual(db.execute("SELECT count(*) FROM occurrences").fetchone()[0],2)
    def test_docx_repair_preserves_history_shared_children_and_idempotence(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,"w") as z:
            z.writestr("[Content_Types].xml","<Types/>"); z.writestr("word/document.xml","<document/>")
        raw=data.getvalue(); sha=hashlib.sha256(raw).hexdigest(); (self.root/"response.docx").write_bytes(raw); (self.root/"shared.xml").write_bytes(b"<document/>")
        self.run_cli("inventory"); db=self.db(); db.execute("UPDATE docs SET format='zip' WHERE sha=?",(sha,)); db.commit(); self.run_cli("extract")
        old_edges=db.execute("SELECT count(*) FROM edges WHERE parent=?",(sha,)).fetchone()[0]; self.assertEqual(old_edges,2)
        db.execute("DELETE FROM extraction_attempts WHERE sha=?",(sha,)); db.commit()
        before=[tuple(r) for r in db.execute("SELECT oid,sha,receipt,first_seen FROM occurrences ORDER BY oid,sha")]
        def repair():
            p=subprocess.run([sys.executable,str(SCRIPT.parent/"repair_intake.py"),"--output",str(self.out)],capture_output=True,text=True)
            self.assertEqual(p.returncode,0,p.stderr); return json.loads(p.stdout.splitlines()[-1])
        result=repair(); self.assertEqual(result["docx_hashes_reclassified"],[sha]); self.assertEqual(result["edges_retired_this_run"],2)
        self.assertEqual(tuple(db.execute("SELECT format,stage FROM docs WHERE sha=?",(sha,)).fetchone()),("docx","complete"))
        self.assertEqual(db.execute("SELECT count(*) FROM edges WHERE parent=?",(sha,)).fetchone()[0],0)
        self.assertEqual(db.execute("SELECT count(*) FROM edge_history WHERE parent=?",(sha,)).fetchone()[0],2)
        prior=json.loads(db.execute("SELECT record FROM document_history WHERE sha=?",(sha,)).fetchone()[0]); self.assertEqual(prior["format"],"zip"); self.assertEqual(prior["stage"],"complete")
        legacy=db.execute("SELECT stage,version,artifact_dir,finished FROM extraction_attempts WHERE sha=? AND id LIKE 'legacy-%'",(sha,)).fetchone()
        self.assertEqual(legacy[0],"complete"); self.assertEqual(legacy[1],prior["version"]); self.assertIsNone(legacy[3]); self.assertTrue(pathlib.Path(legacy[2],"digest.json").exists())
        self.assertEqual(before,[tuple(r) for r in db.execute("SELECT oid,sha,receipt,first_seen FROM occurrences ORDER BY oid,sha")])
        active,_=m.active(db); self.assertEqual(set(active),{sha,hashlib.sha256(b"<document/>").hexdigest()})
        for retired in result["retired_hashes"]: self.assertTrue((self.out/"blobs"/retired).is_file())
        again=repair(); self.assertEqual(again["docx_hashes_reclassified"],[]); self.assertEqual(again["permissions_repaired"],[]); self.assertEqual(again["edges_retired_this_run"],0); self.assertEqual(again["errors"],[])

if __name__=="__main__": unittest.main(verbosity=2)
