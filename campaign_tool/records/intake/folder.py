#!/usr/bin/env python3
"""Read-only evidence intake. Inventory requires only Python's standard library."""
import argparse, collections, csv, datetime, email.policy, hashlib, io, json, os
import pathlib, re, sqlite3, stat, subprocess, sys, time, uuid, zipfile
from email.parser import BytesParser

# Direct script execution, including worker subprocesses, has no package context.
if __package__ in {None, ""}:
    sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[3]))
from campaign_tool.records.config import load_config, validate_config

VERSION = "flock-intake-3.2"
ENGINE_REVISION = "portable-intake-3"
LIMITS = dict(source_bytes=512*1024**2, member_bytes=128*1024**2, expanded_bytes=1024**3,
              members=10000, depth=5, ratio=1000, seconds=180, memory_bytes=3*1024**3)
SKIP_DIRS = {".git", "node_modules", ".venv", "venv", "__pycache__", ".cache", "deployment", "deployments", "site-packages", "dist-packages"}
CODE_EXT = {".py", ".pyc", ".js", ".ts", ".tsx", ".jsx", ".sh", ".bash", ".go", ".rb", ".exe", ".so", ".dylib", ".cjs", ".mjs", ".h", ".hpp", ".hxx", ".c", ".cc", ".cpp", ".cxx", ".a", ".o", ".lib", ".dll", ".typed", ".pyi", ".whl", ".egg"}
SCHEMA = """\nPRAGMA foreign_keys=ON;\nCREATE TABLE IF NOT EXISTS runs(id TEXT PRIMARY KEY,started TEXT,finished TEXT,kind TEXT,status TEXT,detail TEXT);\nCREATE TABLE IF NOT EXISTS docs(sha TEXT PRIMARY KEY,bytes INTEGER,format TEXT,stage TEXT DEFAULT 'pending',\n first_seen TEXT,review_status TEXT DEFAULT 'not_reviewed',version TEXT,digest TEXT);\nCREATE TABLE IF NOT EXISTS occurrences(oid TEXT,sha TEXT REFERENCES docs(sha),path TEXT,root TEXT,parent TEXT,\n locator TEXT,class TEXT,receipt TEXT,first_seen TEXT,last_seen TEXT,PRIMARY KEY(oid,sha));\nCREATE TABLE IF NOT EXISTS seen(run TEXT,oid TEXT,sha TEXT,PRIMARY KEY(run,oid));\nCREATE TABLE IF NOT EXISTS edges(parent TEXT,locator TEXT,child TEXT,name TEXT,PRIMARY KEY(parent,locator));\nCREATE TABLE IF NOT EXISTS units(sha TEXT,ordinal INTEGER,kind TEXT,locator TEXT,text TEXT,data TEXT,\n PRIMARY KEY(sha,ordinal));\nCREATE TABLE IF NOT EXISTS events(run TEXT,stage TEXT,code TEXT,path TEXT,sha TEXT,locator TEXT,detail TEXT);\nCREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY,value TEXT);\nCREATE TABLE IF NOT EXISTS preservations(sha TEXT PRIMARY KEY REFERENCES docs(sha),path TEXT,verified_at TEXT,bytes INTEGER);\nCREATE TABLE IF NOT EXISTS scope_exclusions(path TEXT,sha TEXT,reason TEXT,at TEXT,PRIMARY KEY(path,sha));\nCREATE TABLE IF NOT EXISTS extraction_attempts(id TEXT PRIMARY KEY,sha TEXT,started TEXT,finished TEXT,version TEXT,limits TEXT,stage TEXT,digest TEXT,artifact_dir TEXT);\nCREATE TABLE IF NOT EXISTS document_history(sha TEXT,changed_at TEXT,reason TEXT,record TEXT);\nCREATE TABLE IF NOT EXISTS edge_history(parent TEXT,locator TEXT,child TEXT,name TEXT,invalidated_at TEXT,reason TEXT,source_version TEXT);\nCREATE INDEX IF NOT EXISTS occurrence_sha ON occurrences(sha);\nCREATE INDEX IF NOT EXISTS unit_sha ON units(sha);\n"""
def now(): return datetime.datetime.now(datetime.timezone.utc).isoformat()
def js(x): return json.dumps(x, ensure_ascii=True, sort_keys=True, default=str)
def hid(x): return hashlib.sha256(x.encode()).hexdigest()
def atomic(path, value):
    p = pathlib.Path(path); p.parent.mkdir(parents=True, exist_ok=True)
    t = p.with_name(p.name + ".tmp"); t.write_text(js(value)+"\n"); os.replace(t, p)
def event(db, run, stage, code, path="", sha="", locator=None, detail=""):
    db.execute("INSERT INTO events VALUES(?,?,?,?,?,?,?)", (run,stage,code,path,sha,js(locator),str(detail)))
def excluded(path, out=None, excluded_path_fragments=()):
    p = pathlib.PurePosixPath(str(path).replace("\\", "/")); n = p.name.lower()
    if out and (str(path)==str(out) or str(path).startswith(str(out)+"/")): return "own_output"
    if any(x.lower() in SKIP_DIRS for x in p.parts): return "software_or_cache_directory"
    if any(x.lower().endswith((".dist-info",".egg-info")) for x in p.parts): return "python_package_metadata"
    if any(fragment in str(p) for fragment in excluded_path_fragments): return "configured_path_fragment"
    if any(x.lower().startswith("analysis-pipeline-") for x in p.parts): return "pipeline_output"
    if n.startswith((".env", "credentials", "service-account", "secrets", "id_rsa", "id_ed25519")) or p.suffix.lower() in {".pem", ".key", ".p12", ".pfx"}: return "credential_or_secret_name"
    if n in {"dockerfile", "package-lock.json", "package.json", "yarn.lock", "pnpm-lock.yaml"} or n.startswith("docker-compose"): return "software_configuration"
    if p.suffix.lower() in CODE_EXT: return "executable_or_software_source"
    if n==".ds_store" or n.startswith("._") or "__MACOSX" in p.parts: return "filesystem_metadata"
    return None
def provenance(path):
    s = str(path).lower()
    if "/originals/" in s or "/productions/" in s or "/received/" in s or "/mail-history/" in s: return "original"
    if re.search(r"analysis|deep.dive|crosscheck|inventory|catalog|tracker|draft|handoff|research|audit-|follow.up|review|campaign|launch|cutover|website|outreach|tracking|status|register", pathlib.Path(s).name): return "generated_analysis"
    return "unknown"
def fmt(name, head=b"", container=None):
    ext = pathlib.Path(name).suffix.lower().lstrip(".")
    if b"%PDF-" in head[:1024]: return "pdf"
    if ext in {"xlsx","xlsm","xltx","xltm"}: return "xlsx"
    if ext=="docx": return "docx"
    if head.startswith(b"PK\x03\x04"):
        if container is not None:
            try:
                with (io.BytesIO(container) if isinstance(container,bytes) else secure_open(container)) as f:
                    size=f.seek(0,2); f.seek(max(0,size-65557)); tail=f.read(65557); i=tail.rfind(b"PK\x05\x06"); end=tail[i:i+22]
                    if i>=0 and len(end)==22 and int.from_bytes(end[10:12],"little")<=LIMITS["members"] and int.from_bytes(end[12:16],"little")<=8*1024**2:
                        f.seek(0)
                        with zipfile.ZipFile(f) as z:
                            names=z.namelist()
                            if all(safe_member(n) for n in names) and "xl/workbook.xml" in names: return "xlsx"
                            if all(safe_member(n) for n in names) and {"[Content_Types].xml","word/document.xml"}<=set(names): return "docx"
            except (OSError,ValueError,zipfile.BadZipFile): pass
        return "zip"
    if not ext or ext in {"bin","unknown"}:
        header=re.split(br"\r?\n\r?\n",head[:16384],maxsplit=1)[0]
        keys=set(k.lower() for k in re.findall(br"(?m)^([A-Za-z][A-Za-z0-9-]{0,40}):",header))
        if len(keys & {b"from",b"to",b"subject",b"date",b"message-id",b"mime-version",b"received"})>=3 and keys & {b"from",b"received"}: return "eml"
    return ext or "unknown"
def secure_open(path):
    parts = pathlib.Path(os.path.abspath(path)).parts; fd = os.open("/", os.O_RDONLY|os.O_DIRECTORY)
    try:
        for part in parts[1:-1]:
            nxt = os.open(part, os.O_RDONLY|os.O_DIRECTORY|os.O_NOFOLLOW, dir_fd=fd); os.close(fd); fd=nxt
        f = os.open(parts[-1], os.O_RDONLY|os.O_NOFOLLOW, dir_fd=fd)
        if not stat.S_ISREG(os.fstat(f).st_mode): os.close(f); raise ValueError("not a regular file")
        return os.fdopen(f, "rb")
    finally: os.close(fd)
def check_space(out,size):
    s=os.statvfs(out)
    if s.f_bavail*s.f_frsize<size+64*1024**2: raise OSError("insufficient_disk_for_preservation")
def verify_blob(sha,out,make_readonly=False):
    if not re.fullmatch(r"[a-f0-9]{64}",sha): raise ValueError("invalid_blob_hash")
    with secure_open(out/"blobs"/sha) as f:
        before=os.fstat(f.fileno()); h=hashlib.sha256()
        for block in iter(lambda:f.read(1024**2),b""): h.update(block)
        if h.hexdigest()!=sha: raise ValueError("existing_blob_hash_mismatch; not modified")
        if make_readonly and stat.S_IMODE(before.st_mode)!=0o400:
            if before.st_nlink!=1: raise ValueError("refusing_shared_inode_permission_change")
            os.fchmod(f.fileno(),0o400); os.fsync(f.fileno())
        after=os.fstat(f.fileno())
    return dict(sha256=sha,bytes=after.st_size,before_mode=stat.S_IMODE(before.st_mode),after_mode=stat.S_IMODE(after.st_mode))
def seal_blob(tmp,sha,out):
    dest=out/"blobs"/sha
    try:
        if dest.exists():
            verify_blob(sha,out,make_readonly=True)
        else:
            os.chmod(tmp,0o400); os.link(tmp,dest)
            fd=os.open(dest.parent,os.O_RDONLY|os.O_DIRECTORY)
            try: os.fsync(fd)
            finally: os.close(fd)
        return str(dest)
    finally: pathlib.Path(tmp).unlink(missing_ok=True)
def upgrade_format(db,sha,form):
    old=db.execute("SELECT * FROM docs WHERE sha=?",(sha,)).fetchone()
    priority=lambda f: 0 if f=="unknown" else 1 if f=="zip" else 3 if f in {"pdf","xlsx","eml","docx"} else 2
    if not old or priority(form)<=priority(old["format"]): return False
    stamp=now(); reason="format_upgrade:"+old["format"]+"->"+form
    db.execute("INSERT INTO document_history VALUES(?,?,?,?)",(sha,stamp,reason,js(dict(old))))
    db.execute("INSERT INTO edge_history SELECT parent,locator,child,name,?,?,? FROM edges WHERE parent=?",(stamp,reason,old["version"],sha))
    db.execute("DELETE FROM edges WHERE parent=?",(sha,))
    db.execute("UPDATE docs SET format=?,stage='pending',version=NULL WHERE sha=?",(form,sha))
    return True
def register(db, sha, size, form, path, root, parent=None, locator=None, receipt=None):
    stamp=now(); oid=hid(js([root,path,parent,locator])); cl=provenance(path) if not parent else "unknown"
    db.execute("INSERT OR IGNORE INTO docs(sha,bytes,format,first_seen) VALUES(?,?,?,?)", (sha,size,form,stamp))
    upgrade_format(db,sha,form)
    db.execute("INSERT INTO occurrences VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(oid,sha) DO UPDATE SET last_seen=excluded.last_seen,class=excluded.class",
               (oid,sha,path,root,parent,js(locator),cl,js(receipt or {}),stamp,stamp))
    return oid
def inventory(db, out, roots, config=None):
    config = config if config is not None else load_config()
    run=uuid.uuid4().hex; started=time.monotonic(); counts=collections.Counter(); root_stats=[]
    db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?)", (run,now(),None,"inventory","running",js(roots))); db.commit()
    manifest=out/"manifests"/(run+".jsonl"); manifest.parent.mkdir(exist_ok=True); (out/"blobs").mkdir(exist_ok=True)
    blocked_paths={r[0] for r in db.execute("SELECT path FROM scope_exclusions")}; blocked_hashes={r[0] for r in db.execute("SELECT sha FROM scope_exclusions WHERE sha!=''")}
    with manifest.open("w") as log:
        def record(row): log.write(js(row)+"\n"); log.flush(); counts[row["status"]]+=1
        for root in roots:
            rs=dict(root=root,hashed=0,errors=0,excluded=0); root_stats.append(rs)
            def onerror(e):
                rs["errors"]+=1; record(dict(status="error",root=root,path=e.filename,error=str(e)))
                event(db,run,"inventory","scan_error",e.filename or root,detail=e)
            try:
                if pathlib.Path(root).resolve()!=pathlib.Path(root) or not os.path.isdir(root): raise ValueError("missing root or symlink in root path")
                for base, dirs, files in os.walk(root, topdown=True, followlinks=False, onerror=onerror):
                    for name in sorted(dirs+files):
                        path=os.path.join(base,name)
                        try:
                            st=os.lstat(path); relative=os.path.relpath(path,root)
                            why="out_of_scope_personal" if path in blocked_paths else "symlink" if stat.S_ISLNK(st.st_mode) else ("own_output" if path==str(out) or path.startswith(str(out)+"/") else excluded(relative, excluded_path_fragments=config["excluded_path_fragments"]))
                            if why:
                                if name in dirs: dirs.remove(name)
                                rs["excluded"]+=1; record(dict(status="excluded",root=root,path=path,reason=why,directory=stat.S_ISDIR(st.st_mode)))
                                event(db,run,"inventory",why,path); continue
                            if stat.S_ISDIR(st.st_mode): continue
                            if not stat.S_ISREG(st.st_mode): raise ValueError("non-regular filesystem object")
                            h=hashlib.sha256(); head=b""; size=0; check_space(out,st.st_size); tmp=out/"blobs"/(".pending-"+uuid.uuid4().hex)
                            try:
                                with secure_open(path) as f, tmp.open("xb") as saved:
                                    before=os.fstat(f.fileno())
                                    while True:
                                        b=f.read(1024**2)
                                        if not b: break
                                        if not head: head=b[:16384]
                                        h.update(b); size+=len(b); saved.write(b)
                                    after=os.fstat(f.fileno()); saved.flush(); os.fsync(saved.fileno())
                                if (before.st_size,before.st_mtime_ns,before.st_ctime_ns)!=(after.st_size,after.st_mtime_ns,after.st_ctime_ns) or size!=after.st_size: raise ValueError("source changed while hashing")
                                if h.hexdigest() in blocked_hashes:
                                    rs["excluded"]+=1; record(dict(status="excluded",path=path,reason="out_of_scope_hash_alias")); event(db,run,"inventory","out_of_scope_hash_alias",path); continue
                                stored=seal_blob(tmp,h.hexdigest(),out)
                            finally: tmp.unlink(missing_ok=True)
                            if (before.st_size,before.st_mtime_ns,before.st_ctime_ns)!=(after.st_size,after.st_mtime_ns,after.st_ctime_ns) or size!=after.st_size: raise ValueError("source changed while hashing")
                            sha=h.hexdigest(); receipt=dict(mtime_ns=before.st_mtime_ns,ctime_ns=before.st_ctime_ns,device=before.st_dev,inode=before.st_ino,observed_at=now(),received_at=None)
                            oid=register(db,sha,size,fmt(path,head,stored),path,root,receipt=receipt)
                            db.execute("INSERT OR IGNORE INTO preservations VALUES(?,?,?,?)",(sha,stored,now(),size))
                            db.execute("INSERT INTO seen VALUES(?,?,?)",(run,oid,sha)); rs["hashed"]+=1
                            record(dict(status="included",oid=oid,sha256=sha,bytes=size,path=path,root=root,classification=provenance(path),classification_basis="trusted directory first; otherwise provisional filename heuristic, never excluded from review",receipt=receipt,storage="preserved",blob_path=stored,name_review_required=not pathlib.Path(path).suffix))
                            if rs["hashed"]%100==0: db.commit()
                        except Exception as e: onerror(OSError(str(e),path))
            except Exception as e:
                rs["errors"]+=1; record(dict(status="error",root=root,path=root,error=str(e))); event(db,run,"inventory","root_error",root,detail=e)
    db.execute("INSERT OR REPLACE INTO meta VALUES('inventory_run',?)",(run,))
    stats=dict(run_id=run,version=VERSION,roots=root_stats,counts=dict(counts),seconds=round(time.monotonic()-started,3),manifest=str(manifest),complete=not counts["error"])
    stats["unique_hashes"]=db.execute("SELECT count(DISTINCT sha) FROM seen WHERE run=?",(run,)).fetchone()[0]
    stats["duplicate_occurrences"]=counts["included"]-stats["unique_hashes"]
    stats["occurrence_bytes"]=db.execute("SELECT coalesce(sum(d.bytes),0) FROM seen s JOIN docs d ON d.sha=s.sha WHERE s.run=?",(run,)).fetchone()[0]
    stats["unique_bytes"]=db.execute("SELECT coalesce(sum(bytes),0) FROM docs WHERE sha IN (SELECT sha FROM seen WHERE run=?)",(run,)).fetchone()[0]
    db.execute("UPDATE runs SET finished=?,status=?,detail=? WHERE id=?",(now(),"complete" if stats["complete"] else "partial",js(stats),run)); db.commit()
    atomic(out/"inventory-summary.json",stats); atomic(out/"manifests"/(run+"-summary.json"),stats); report(db,out,config); print(js(stats),flush=True); return run
def presence(v):
    if v is None: return "missing_field"
    if v=="": return "empty"
    s=str(v)
    if not s.strip(): return "whitespace"
    if re.fullmatch(r"(?i)\s*(?:\[?redacted\]?|\[?withheld\]?|x{3,}|\u2588+)\s*",s): return "redacted_marker"
    return "present"
def safe_member(name):
    p=pathlib.PurePosixPath(name.replace("\\","/"))
    return bool(name) and not p.is_absolute() and ".." not in p.parts and not re.match(r"^[A-Za-z]:",name) and "\x00" not in name
def worker(src, sha, form, dest, out, limits, depth, config=None):
    config = config if config is not None else load_config()
    import resource
    resource.setrlimit(resource.RLIMIT_AS,(limits["memory_bytes"],limits["memory_bytes"]))
    dest.mkdir(parents=True,exist_ok=True); stats=collections.Counter(); issues=[]; children=[]; stage="complete"
    blocked=json.loads((out/"scope-excluded-hashes.json").read_text()) if (out/"scope-excluded-hashes.json").exists() else []
    uf=(dest/"units.jsonl").open("w",buffering=1); tf=(dest/"text.txt").open("w"); started=time.monotonic()
    def issue(code,loc=None,detail=""):
        issues.append(dict(code=code,locator=loc,detail=str(detail)))
    def emit(kind,loc,text="",data=None):
        stats["units"]+=1; stats[kind]+=1
        uf.write(js(dict(kind=kind,locator=loc,text=text,data=data or {}))+"\n")
        tf.write(js(loc)+"\n"+text+"\n")
    def child(name,data,loc,relation):
        if depth>=limits["depth"]: issue("nesting_limit",loc,name); return
        why=excluded(name, excluded_path_fragments=config["excluded_path_fragments"])
        if why: issue("excluded_"+why,loc,name); return
        if len(data)>limits["member_bytes"] or stats["expanded_bytes"]+len(data)>limits["expanded_bytes"] or len(children)>=limits["members"]: issue("expansion_limit",loc,name); return
        h=hashlib.sha256(data).hexdigest()
        if h in blocked: issue("out_of_scope_child",loc); return
        p=out/"blobs"/h; p.parent.mkdir(exist_ok=True); check_space(out,len(data)); tmp=p.parent/(".pending-"+uuid.uuid4().hex)
        with tmp.open("xb") as f: f.write(data); f.flush(); os.fsync(f.fileno())
        seal_blob(tmp,h,out)
        stats["expanded_bytes"]+=len(data); children.append(dict(sha=h,bytes=len(data),name=name,locator=loc,relation=relation,format=fmt(name,data[:16384],data)))
    try:
        with secure_open(src) as f:
            if os.fstat(f.fileno()).st_size>limits["source_bytes"]: raise ValueError("source_size_limit; hashed but extraction queued")
            raw=f.read(limits["source_bytes"]+1)
        if hashlib.sha256(raw).hexdigest()!=sha: raise ValueError("source_hash_mismatch; re-inventory required")
        if form=="pdf":
            from pypdf import PdfReader
            pdf=PdfReader(io.BytesIO(raw),strict=False)
            if pdf.is_encrypted:
                if not pdf.decrypt(""): raise ValueError("encrypted_pdf_requires_user_password; no credentials attempted")
                stats["standard_passwordless_open"]=1
            stats["pages_expected"]=len(pdf.pages)
            for n,page in enumerate(pdf.pages,1):
                loc=dict(page=n)
                try:
                    text=page.extract_text() or ""; needed=not text.strip()
                    if needed: issue("ocr_needed_or_blank_page",loc); stats["ocr_needed"]+=1
                    emit("pdf_page",loc,text,dict(ocr_needed=needed,visual_review="not_done"))
                except Exception as e: issue("page_error",loc,e); emit("pdf_page",loc,"",dict(error=str(e),ocr_needed=True))
        elif form=="eml":
            msg=BytesParser(policy=email.policy.default).parsebytes(raw)
            emit("email_headers",{"mime":"0"},data=dict(headers=list(msg.raw_items()),defects=[str(x) for x in msg.defects]))
            def mime(part,loc):
                if part.defects: issue("mime_defect",loc,[str(x) for x in part.defects])
                if part.get_content_type()=="message/rfc822":
                    for n,m in enumerate(part.get_payload() if isinstance(part.get_payload(),list) else []):
                        child(part.get_filename() or "attached.eml",m.as_bytes(),dict(mime=loc,message=n),"reserialized_rfc822")
                    issue("rfc822_reserialized",loc,"child hash is serialized parsed message, not asserted original attachment octets"); return
                if part.is_multipart():
                    for n,p in enumerate(part.iter_parts(),1): mime(p,loc+"."+str(n))
                    return
                data=part.get_payload(decode=True) or b""; name=part.get_filename(); typ=part.get_content_type()
                emit("mime_part",{"mime":loc},data=dict(content_type=typ,filename=name,disposition=part.get_content_disposition(),headers=list(part.raw_items())))
                if name or part.get_content_disposition()=="attachment" or not typ.startswith("text/"): child(name or ("part-"+loc),data,{"mime":loc},"decoded_mime_payload")
                else:
                    charset=part.get_content_charset() or "utf-8"
                    try: body=data.decode(charset)
                    except (UnicodeError,LookupError): body=data.decode("utf-8",errors="replace"); issue("body_decode_replacement",loc,charset)
                    emit("email_body",{"mime":loc},body,dict(content_type=typ,charset=charset))
            mime(msg,"1")
        elif form in {"csv","tsv"}:
            try: text=raw.decode("utf-8-sig"); encoding="utf-8-sig"
            except UnicodeError:
                text=raw.decode("utf-16") if raw[:2] in {b"\xff\xfe",b"\xfe\xff"} else raw.decode("cp1252"); encoding="utf-16" if raw[:2] in {b"\xff\xfe",b"\xfe\xff"} else "cp1252"
                issue("encoding_fallback",detail=encoding)
            csv.field_size_limit(min(sys.maxsize,limits["source_bytes"])); reader=csv.reader(io.StringIO(text,newline=""),delimiter="\t" if form=="tsv" else ",")
            headers=next(reader,[]); emit("csv_header",dict(record=1,line_start=1,line_end=reader.line_num),js(headers)); fields=[collections.Counter() for _ in headers]; duplicates={}; end=reader.line_num
            for rowno,row in enumerate(reader,2):
                loc=dict(record=rowno,line_start=end+1,line_end=reader.line_num); end=reader.line_num
                states=[presence(row[i] if i<len(row) else None) for i in range(len(headers))]
                for i,s in enumerate(states): fields[i][s]+=1
                key=hid(js(row)); duplicate=duplicates.get(key); duplicates.setdefault(key,rowno)
                if duplicate is not None: stats["duplicate_rows"]+=1
                if len(row)!=len(headers): stats["ragged_rows"]+=1
                emit("csv_row",loc,js(row),dict(values=row,presence=states,duplicate_of_record=duplicate,row_sha256=key))
            stats["fields"]=[dict(index=i,header=h,counts=dict(fields[i])) for i,h in enumerate(headers)]; stats["encoding"]=encoding
        elif form=="docx":
            import xml.etree.ElementTree as ET
            ns="{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                infos=z.infolist()
                if len(infos)>limits["members"] or sum(x.file_size for x in infos)>limits["expanded_bytes"] or any(not safe_member(x.filename) or x.file_size>limits["member_bytes"] or x.file_size/max(x.compress_size,1)>limits["ratio"] for x in infos):
                    raise ValueError("unsafe_or_oversized_docx_zip")
                names=[x.filename for x in infos]
                parts=[n for n in names if n=="word/document.xml" or (n.startswith(("word/header","word/footer")) and n.endswith(".xml")) or n in ("word/footnotes.xml","word/endnotes.xml","word/comments.xml")]
                if "word/document.xml" not in parts: raise ValueError("docx_missing_document_xml")
                for name in parts:
                    root=ET.fromstring(z.read(name))
                    for count,para in enumerate(root.iter(ns+"p"),1):
                        tokens=[]
                        for node in para.iter():
                            if node.tag==ns+"t" and node.text: tokens.append(node.text)
                            elif node.tag==ns+"tab": tokens.append("\\t")
                            elif node.tag in (ns+"br",ns+"cr"): tokens.append("\\n")
                        value="".join(tokens)
                        if value:
                            emit("docx_paragraph",dict(part=name,paragraph=count),value)
                    stats["docx_parts"]+=1
                if any(n.startswith("word/media/") or n.startswith("word/embeddings/") for n in names):
                    issue("docx_media_or_embeddings_not_visually_reviewed")
        elif form=="xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                infos=z.infolist()
                if len(infos)>limits["members"] or sum(x.file_size for x in infos)>limits["expanded_bytes"] or any(not safe_member(x.filename) or x.file_size>limits["member_bytes"] or x.file_size/max(x.compress_size,1)>limits["ratio"] for x in infos): raise ValueError("unsafe_or_oversized_workbook_zip")
                if any("vbaProject" in x.filename or "/embeddings/" in x.filename for x in infos): issue("workbook_embedded_content_not_extracted")
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            for sheet in wb.worksheets:
                emit("workbook_sheet",dict(sheet=sheet.title),data=dict(state=sheet.sheet_state,declared_rows=sheet.max_row,declared_columns=sheet.max_column))
                sheet.reset_dimensions()
                for n,row in enumerate(sheet.iter_rows(),1):
                    cells=[dict(cell=c.coordinate,value=c.value,type=c.data_type) for c in row if c.value is not None]
                    stats["formulas"]+=sum(c["type"]=="f" for c in cells)
                    emit("workbook_row",dict(sheet=sheet.title,row=n),js(cells),dict(cells=cells))
            wb.close(); issue("workbook_cached_values_and_objects_not_reviewed",detail="Formula text retained; no recalculation, macro execution, object rendering or cached-value comparison.")
        elif form=="zip":
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                for n,info in enumerate(z.infolist(),1):
                    loc=dict(member_index=n,member=info.filename)
                    emit("zip_member",loc,data=dict(bytes=info.file_size,compressed_bytes=info.compress_size,crc=info.CRC))
                    if info.is_dir(): continue
                    if n>limits["members"]: issue("zip_member_count_limit",loc); continue
                    if not safe_member(info.filename) or stat.S_ISLNK(info.external_attr>>16): issue("unsafe_zip_member",loc); continue
                    why=excluded(info.filename, excluded_path_fragments=config["excluded_path_fragments"])
                    if why: issue("excluded_"+why,loc); continue
                    if info.flag_bits&1: issue("encrypted_zip_member",loc); continue
                    if info.file_size>limits["member_bytes"] or stats["expanded_bytes"]+info.file_size>limits["expanded_bytes"] or info.file_size/max(info.compress_size,1)>limits["ratio"]: issue("zip_expansion_limit",loc); continue
                    try:
                        with z.open(info) as f: data=f.read(limits["member_bytes"]+1)
                        child(info.filename,data,loc,"zip_member")
                    except Exception as e: issue("zip_member_error",loc,e)
        elif form in {"txt","md","json","jsonl","xml","html","htm","log","rst","yaml","yml"}:
            text=raw.decode("utf-8",errors="replace")
            if "\ufffd" in text: issue("text_decode_replacement")
            for n,line in enumerate(text.splitlines(),1): emit("text_line",dict(line=n),line)
        else: stage="unsupported"; issue("unsupported_format",detail=form)
    except Exception as e: stage="failed"; issue("extraction_error",detail=repr(e))
    finally: uf.close(); tf.close()
    if issues and stage=="complete": stage="partial"
    result=dict(sha256=sha,version=VERSION,stage=stage,review_status="not_reviewed",effective_limits=limits,counts=dict(stats),issues=issues,children=children,seconds=round(time.monotonic()-started,3))
    result["engine_revision"]=ENGINE_REVISION
    result["extraction_policy_sha256"]=extraction_policy(form,config)
    result["children_inventory_complete"]=stage in {"complete","partial"} and all(e["code"].startswith("excluded_") or e["code"]=="out_of_scope_child" for e in issues)
    result["artifacts"]={n:hashlib.sha256((dest/n).read_bytes()).hexdigest() for n in ("units.jsonl","text.txt")}
    atomic(dest/"digest.json",result)
def active(db):
    r=db.execute("SELECT value FROM meta WHERE key='inventory_run'").fetchone()
    if not r: return {},None
    blocked={x[0] for x in db.execute("SELECT sha FROM scope_exclusions")}
    levels={x[0]:0 for x in db.execute("SELECT DISTINCT sha FROM seen WHERE run=?",(r[0],)) if x[0] not in blocked}
    todo=list(levels)
    for p in todo:
        for (child,) in db.execute("SELECT child FROM edges WHERE parent=?",(p,)):
            if child not in levels and child not in blocked: levels[child]=levels[p]+1; todo.append(child)
    return levels,r[0]
def extraction_policy(form,config):
    """Fingerprint parser-affecting settings, not host paths or agency routing."""
    return hid(js(dict(format=form,paragraph_ordinals=2 if form=="docx" else None,
                       excluded_path_fragments=sorted(set(config["excluded_path_fragments"])) if form in {"zip","eml"} else [])))
def extraction_budget(form,limits,depth):
    """Bind effective parser resources, not the container's absolute location."""
    return dict(limits={key:limits[key] for key in LIMITS if key!="depth"},
                remaining_depth=max(0,limits["depth"]-depth) if form in {"zip","eml"} else None)
def reconcile_edges(db,sha,digest,config,source_version):
    """Retire only proven obsolete relationships; retain immutable history.

    Partial/failed enumeration does not prove an unseen child absent. Explicit
    exclusions still retire matching links. Other parents/direct originals keep
    shared child hashes active; old blobs, receipts and units are not deleted.
    """
    from .rfc822_adapter import SCHEMA as wire_schema
    incoming={js(child["locator"]):child for child in digest.get("children",[])}
    complete=digest.get("children_inventory_complete") is True
    stamp=now()
    for edge in list(db.execute("SELECT * FROM edges WHERE parent=?",(sha,))):
        # A legacy extraction inventory cannot prove exact-wire acquisition absent.
        try:
            locator=json.loads(edge["locator"])
        except (TypeError,ValueError):
            locator=None
        wire_acquisition=(isinstance(locator,dict) and set(locator)=={"mime","wire_schema"}
                          and locator["wire_schema"]==wire_schema
                          and isinstance(locator["mime"],str)
                          and len(locator["mime"])<=256
                          and re.fullmatch(r"1(?:\.[1-9][0-9]*)*",locator["mime"]) is not None)
        new=incoming.get(edge["locator"])
        reason=None
        if new and (new["sha"]!=edge["child"] or new["name"]!=edge["name"]):
            reason="observed_relationship_replaced"
        elif not new and excluded(edge["name"],excluded_path_fragments=config["excluded_path_fragments"]):
            reason="relationship_excluded_by_current_policy"
        elif not new and complete and not wire_acquisition:
            reason="absent_from_completed_child_inventory"
        if reason:
            db.execute("INSERT INTO edge_history VALUES(?,?,?,?,?,?,?)",
                       (sha,edge["locator"],edge["child"],edge["name"],stamp,reason,source_version))
            db.execute("DELETE FROM edges WHERE parent=? AND locator=?",(sha,edge["locator"]))
def extract(db,out,limits,retry=False,selection=None,config=None):
    config = config if config is not None else load_config(output=str(out))
    levels,inv=active(db)
    if selection is not None and (not re.fullmatch(r"[a-f0-9]{64}",selection) or selection not in levels): raise ValueError("--sha must be a known active lowercase SHA-256; invalid/inactive selector rejected")
    run=uuid.uuid4().hex; started=time.monotonic(); db.execute("INSERT INTO runs VALUES(?,?,?,?,?,?)",(run,now(),None,"extraction","running",js(dict(limits=limits,selected_sha=selection)))); db.commit(); done=0
    todo=[selection] if selection else list(sorted(levels,key=lambda x:(levels[x],x)))
    queued=set(todo)
    for sha in todo:
        if sha not in levels: continue
        doc=dict(db.execute("SELECT * FROM docs WHERE sha=?",(sha,)).fetchone())
        prior_digest=json.loads(doc["digest"] or "{}")
        policy=extraction_policy(doc["format"],config)
        previous_policy=prior_digest.get("extraction_policy_sha256")
        invocation_depth=levels[sha]
        budget=extraction_budget(doc["format"],limits,invocation_depth)
        budget_hash=hid(js(budget))
        previous_budget=prior_digest.get("extraction_budget_sha256")
        # Legacy container depth cannot be inferred from today's occurrence graph.
        # Unchanged non-container work can retain its recorded effective limits.
        if previous_budget is None and doc["format"] not in {"zip","eml"}:
            old_limits=prior_digest.get("effective_limits")
            if isinstance(old_limits,dict) and all(key in old_limits for key in LIMITS):
                previous_budget=hid(js(extraction_budget(doc["format"],old_limits,0)))
        # Old non-container/non-DOCX parsers are unchanged: keep finished work.
        if previous_policy is None and doc["format"] not in {"zip","eml","docx"}:
            previous_policy=policy
        if doc["version"]==VERSION and previous_policy==policy and previous_budget==budget_hash and doc["stage"] not in {"pending","running"} and not retry: continue
        sources=db.execute("SELECT o.* FROM occurrences o JOIN seen s ON o.oid=s.oid AND o.sha=s.sha WHERE s.run=? AND o.sha=?",(inv,sha)).fetchall()
        blob=out/"blobs"/sha; src=str(blob) if blob.exists() else next((x["path"] for x in sources if os.path.isfile(x["path"])),str(blob)); dest=out/"derived"/sha
        attempt=uuid.uuid4().hex
        if dest.exists():
            previous=out/"history"/sha/attempt; previous.parent.mkdir(parents=True,exist_ok=True); os.rename(dest,previous)
            changed=db.execute("UPDATE extraction_attempts SET artifact_dir=? WHERE sha=? AND artifact_dir=?",(str(previous),sha,str(dest))).rowcount
            if not changed and doc["digest"]:
                prior=json.loads(doc["digest"])
                db.execute("INSERT INTO extraction_attempts VALUES(?,?,?,?,?,?,?,?,?)",("legacy-"+attempt,sha,None,None,prior.get("version",doc["version"]),js(prior["effective_limits"]) if "effective_limits" in prior else None,prior.get("stage",doc["stage"]),doc["digest"],str(previous)))
        dest.mkdir(parents=True,exist_ok=True)
        db.execute("INSERT INTO extraction_attempts VALUES(?,?,?,?,?,?,?,?,?)",(attempt,sha,now(),None,VERSION,js(limits),"running",None,str(dest)))
        db.execute("UPDATE docs SET stage='running' WHERE sha=?",(sha,)); db.commit()
        cmd=[sys.executable,__file__,"_worker","--output",str(out),"--source",src,"--sha",sha,"--format",doc["format"],"--depth",str(levels[sha]),"--limits",js(limits),"--worker-config",js({**config,"output":str(out)})]
        with (dest/"stderr.log").open("w") as err:
            try:
                p=subprocess.run(cmd,stdout=err,stderr=err,timeout=limits["seconds"])
                failure="worker_exit_"+str(p.returncode) if p.returncode else "missing_worker_digest"
            except subprocess.TimeoutExpired: failure="worker_timeout"
        digest=json.loads((dest/"digest.json").read_text()) if (dest/"digest.json").exists() else dict(stage="failed",counts={},issues=[dict(code=failure)],children=[],version=VERSION,review_status="not_reviewed")
        digest["engine_revision"]=ENGINE_REVISION
        digest["extraction_policy_sha256"]=policy
        digest["extraction_budget"]=budget
        digest["extraction_budget_sha256"]=budget_hash
        digest["invocation_depth"]=invocation_depth
        digest["effective_limits"]=dict(limits)
        if failure!="missing_worker_digest":
            digest["children_inventory_complete"]=False
            digest["stage"]="failed"
            digest["issues"].append(dict(code=failure))
        db.execute("DELETE FROM units WHERE sha=?",(sha,))
        if (dest/"units.jsonl").exists():
            with (dest/"units.jsonl").open() as f:
                for n,line in enumerate(f,1):
                    try: u=json.loads(line)
                    except ValueError: digest["issues"].append(dict(code="truncated_unit",line=n)); digest["stage"]="partial"; digest["children_inventory_complete"]=False; break
                    db.execute("INSERT INTO units VALUES(?,?,?,?,?,?)",(sha,n,u["kind"],js(u["locator"]),u["text"],js(u["data"])))
        reconcile_edges(db,sha,digest,config,doc["version"])
        for c in digest.get("children",[]):
            loc=js(c["locator"]); register(db,c["sha"],c["bytes"],c["format"],c["name"],"container",sha,c["locator"],dict(relation=c["relation"],parent_sha256=sha))
            db.execute("INSERT OR IGNORE INTO preservations VALUES(?,?,?,?)",(c["sha"],str(out/"blobs"/c["sha"]),now(),c["bytes"]))
            db.execute("INSERT OR REPLACE INTO edges VALUES(?,?,?,?)",(sha,loc,c["sha"],c["name"]))
        levels,_=active(db)
        if selection is None:
            for child in sorted(levels,key=lambda x:(levels[x],x)):
                if child not in queued: todo.append(child); queued.add(child)
        atomic(dest/"digest.json",digest)
        for e in digest["issues"]: event(db,run,"extraction",e["code"],src,sha,e.get("locator"),e.get("detail",""))
        db.execute("UPDATE docs SET stage=?,version=?,digest=? WHERE sha=?",(digest["stage"],VERSION,js(digest),sha))
        db.execute("UPDATE extraction_attempts SET finished=?,stage=?,digest=? WHERE id=?",(now(),digest["stage"],js(digest),attempt)); db.commit(); done+=1
        if done%25==0: print(js(dict(extracted_this_run=done,discovered=len(todo),seconds=round(time.monotonic()-started,1))),flush=True)
    summary=dict(processed=done,selected_sha=selection,seconds=round(time.monotonic()-started,3),run_id=run)
    db.execute("UPDATE runs SET finished=?,status='finished',detail=? WHERE id=?",(now(),js(summary),run)); db.commit(); report(db,out,config); print(js(summary),flush=True)
def report(db,out,config=None):
    config = config if config is not None else load_config()
    levels,inv=active(db); docs={r["sha"]:dict(r) for r in db.execute("SELECT * FROM docs") if r["sha"] in levels}
    paths=collections.defaultdict(list); classes=collections.defaultdict(set); agents=collections.defaultdict(set); stored={r[0] for r in db.execute("SELECT sha FROM preservations")}
    for r in db.execute("SELECT o.* FROM occurrences o JOIN seen s ON o.oid=s.oid AND o.sha=s.sha WHERE s.run=?",(inv,)):
        if r["sha"] not in levels: continue
        paths[r["sha"]].append(r["path"]); classes[r["sha"]].add(r["class"])
        agents[r["sha"]].update(a for a,pattern in config["agency_patterns"].items() if re.search(pattern,r["path"],re.I))
    physical_classes=collections.Counter("original" if "original" in v else "generated_analysis" if "generated_analysis" in v else "unknown" for v in classes.values())
    edges=[dict(r) for r in db.execute("SELECT * FROM edges")]
    for _ in range(LIMITS["depth"]+1):
        for e in edges:
            if e["parent"] in levels:
                agents[e["child"]].update(agents[e["parent"]]); classes[e["child"]].update(classes[e["parent"]])
    stages=collections.Counter(); types=collections.Counter(); unit_counts=collections.Counter(); issue_counts=collections.Counter(); agency_counts=collections.Counter()
    with (out/"all-agency-document-queue.jsonl").open("w") as q:
        for sha,d in sorted(docs.items()):
            stages[d["stage"]]+=1; types[d["format"]]+=1; digest=json.loads(d["digest"] or "{}"); issues=digest.get("issues",[])
            issue_counts.update(e["code"] for e in issues); hints=sorted(agents[sha]) or ["UNASSIGNED"]; agency_counts.update(hints)
            q.write(js(dict(document_id="sha256:"+sha,sha256=sha,bytes=d["bytes"],format=d["format"],stage=d["stage"],extraction_version=d["version"],review_status=d["review_status"],storage="preserved" if sha in stored else "indexed_only",blob_path=str(out/"blobs"/sha) if sha in stored else None,name_review_required=any(not pathlib.Path(p).suffix for p in paths[sha]) or d["format"]=="unknown",agency_hints=hints,agency_basis="path routing only, not confirmed custodian",classification=sorted(classes[sha]) or ["unknown"],source_paths=paths[sha],parents=[e for e in edges if e["child"]==sha],digest_path=str(out/"derived"/sha/"digest.json") if d["digest"] else None,issues=issues))+"\n")
    for r in db.execute("SELECT sha,kind,count(*) AS n FROM units GROUP BY sha,kind"):
        if r["sha"] in levels: unit_counts[r["kind"]]+=r["n"]
    with (out/"occurrences.jsonl").open("w") as f:
        for r in db.execute("SELECT * FROM occurrences ORDER BY oid,first_seen"): f.write(js(dict(r))+"\n")
    with (out/"events.jsonl").open("w") as f:
        for r in db.execute("SELECT * FROM events"): f.write(js(dict(r))+"\n")
    summary=dict(version=VERSION,inventory_run=inv,active_unique_documents=len(docs),physical_unique_hashes=db.execute("SELECT count(DISTINCT sha) FROM seen WHERE run=?",(inv,)).fetchone()[0],physical_occurrences=db.execute("SELECT count(*) FROM seen WHERE run=?",(inv,)).fetchone()[0],stages=dict(stages),formats=dict(types),units=dict(unit_counts),issues=dict(issue_counts),agency_hint_counts=dict(agency_counts),semantic_reviews_complete=0,legal_comparisons_complete=0,independent_reviews_complete=0,mailbox_coverage="not established; filesystem snapshot only",limitations=["Path classifications and agencies are heuristics; prior analyses are leads, not primary proof.","No OCR or visual review performed; text-bearing PDF pages may still contain unread image regions.","CSV values retained without deduplicating rows. Workbook formulas not executed.","Archive limits apply per container and depth; excluded directory descendants are not enumerated.","Historical hashes and receipt metadata retained; physical source bytes are not copied.","No factual, legal-conflict, missing-document or redaction conclusions are auto-approved."])
    summary.update(physical_unique_by_class=dict(physical_classes),active_original_hashes=sum("original" in classes[s] for s in docs),preserved_active_hashes=sum(s in stored for s in docs),indexed_only_active_hashes=sum(s not in stored for s in docs),extraction_version_mismatches=sum(d["version"] not in {None,VERSION} for d in docs.values()))
    summary.update(physical_unique_hashes=sum(v==0 for v in levels.values()),physical_occurrences=db.execute("SELECT count(*) FROM seen WHERE run=? AND sha NOT IN (SELECT sha FROM scope_exclusions)",(inv,)).fetchone()[0],quarantined_hashes=db.execute("SELECT count(DISTINCT sha) FROM scope_exclusions WHERE sha!=''").fetchone()[0])
    summary.update(retired_container_edges=db.execute("SELECT count(*) FROM edge_history").fetchone()[0],format_history_records=db.execute("SELECT count(*) FROM document_history").fetchone()[0])
    summary["limitations"][4]="Content-addressed physical originals and children are preserved read-only from v2 onward; older hash-only history may lack bytes."
    atomic(out/"coverage.json",summary); atomic(out/"reports"/((inv or "none")+"-"+uuid.uuid4().hex+".json"),summary)
    (out/"coverage.md").write_text("# Flock intake coverage\n\n```json\n"+json.dumps(summary,indent=2)+"\n```\n")
def main():
    p=argparse.ArgumentParser(description=__doc__); p.add_argument("command",choices=["inventory","extract","run","report","_worker"])
    p.add_argument("--config",help="explicit private JSON config; CLI roots/output override its paths")
    p.add_argument("--worker-config",help=argparse.SUPPRESS)
    p.add_argument("--output"); p.add_argument("--root",action="append"); p.add_argument("--retry",action="store_true")
    p.add_argument("--source"); p.add_argument("--sha",help="extract: restrict to exactly one known active hash; combine with --retry to re-extract"); p.add_argument("--format"); p.add_argument("--depth",type=int,default=0); p.add_argument("--limits",default="{}")
    args=p.parse_args()
    if args.worker_config is not None and args.command!="_worker": p.error("--worker-config is internal to _worker")
    try:
        config=load_config(args.config,roots=args.root,output=args.output,snapshot=args.worker_config)
        validate_config(config,args.command)
        limits={**LIMITS,**json.loads(args.limits)}
    except (OSError,ValueError,TypeError) as error: p.error(str(error))
    if args.command=="extract" and args.limits!="{}" and not args.sha: p.error("extract limit overrides require --sha for a bounded per-document retry")
    out=pathlib.Path(config["output"]); os.umask(0o077); out.mkdir(parents=True,exist_ok=True)
    if args.command=="_worker": worker(args.source,args.sha,args.format,out/"derived"/args.sha,out,limits,args.depth,config); return
    import fcntl
    with (out/"writer.lock").open("a") as lock:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        db=sqlite3.connect(out/"intake.sqlite"); db.row_factory=sqlite3.Row; db.executescript(SCHEMA); db.execute("PRAGMA journal_mode=WAL")
        try:
            if args.command in {"inventory","run"}: inventory(db,out,config["roots"],config)
            if args.command in {"extract","run"}: extract(db,out,limits,args.retry,args.sha,config)
            if args.command=="report": report(db,out,config)
        finally: db.close()
if __name__=="__main__": main()
